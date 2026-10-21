import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def processar_dados():
    arquivo_caminhos = '/home/kleyfferson/Documentos/projetos_mestrado/projetos-TT-Jo-o/FlapyRepos/caminhos.txt'
    
    if not os.path.exists(arquivo_caminhos):
        print(f"Erro: O arquivo {arquivo_caminhos} não foi encontrado.")
        return

    with open(arquivo_caminhos, 'r', encoding='utf-8') as f:
        caminhos = [linha.strip() for linha in f if '/Test-' in linha]

    print(f"Processando {len(caminhos)} diretórios. Aplicando correção rigorosa de Desvio Padrão...")

    resultados = []

    for path in caminhos:
        if not os.path.exists(path):
            continue

        summary_file = os.path.join(path, 'summary_coverage.csv')
        if not os.path.exists(summary_file):
            continue

        partes = [p for p in path.split('/') if p]
        hash_val = partes[-1]
        projeto_val = partes[-2].replace('Test-', '')

        try:
            df = pd.read_csv(summary_file)
            df.columns = df.columns.str.strip()

            col_resultado = 'Test Result'
            col_cobertura = 'Percentual de Cobertura (%)'

            if col_resultado not in df.columns or col_cobertura not in df.columns:
                continue

            # Padroniza resultado para caixa baixa
            df[col_resultado] = df[col_resultado].astype(str).str.strip().str.lower()
            
            # Limpa qualquer vírgula ou % para ter o número limpo
            df[col_cobertura] = df[col_cobertura].astype(str).str.replace(',', '.', regex=False)
            df[col_cobertura] = df[col_cobertura].str.replace('%', '', regex=False).str.strip()
            
            # Converte para numérico e JOGA FORA dados que não são números (NaNs)
            df[col_cobertura] = pd.to_numeric(df[col_cobertura], errors='coerce')

            # Isola os valores já removendo células vazias (dropna)
            series_passed = df[df[col_resultado] == 'passed'][col_cobertura].dropna()
            series_failed = df[df[col_resultado] == 'failed'][col_cobertura].dropna()

            # --- LÓGICA BLINDADA DO PASSED ---
            if len(series_passed) == 0:
                avg_p, sd_p = 0.0, 0.0
            else:
                avg_p = series_passed.mean()
                # A TRAVA MATEMÁTICA: Se todos os valores da lista forem iguais, o desvio é ZERO.
                if series_passed.nunique() <= 1:
                    sd_p = 0.0
                else:
                    sd_p = series_passed.std(ddof=0)

            # --- LÓGICA BLINDADA DO FAILED ---
            if len(series_failed) == 0:
                avg_f, sd_f = 0.0, 0.0
            else:
                avg_f = series_failed.mean()
                # A TRAVA MATEMÁTICA: Se todos os valores da lista forem iguais, o desvio é ZERO.
                if series_failed.nunique() <= 1:
                    sd_f = 0.0
                else:
                    sd_f = series_failed.std(ddof=0)

            # Arredondar para 6 casas decimais e matar resíduos de ponto flutuante
            avg_p = round(float(avg_p), 6)
            sd_p = round(float(sd_p), 6)
            avg_f = round(float(avg_f), 6)
            sd_f = round(float(sd_f), 6)

            resultados.append({
                'Projeto': projeto_val,
                'Hash': hash_val,
                'AVG_P': avg_p,
                'SD_P': sd_p,
                'AVG_F': avg_f,
                'SD_F': sd_f
            })

        except Exception as e:
            print(f"[Erro] Falha em {projeto_val} ({hash_val}): {e}")

    resultados = sorted(resultados, key=lambda x: x['Projeto'].lower())
    gerar_excel(resultados)


def gerar_excel(resultados):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados Estatisticos"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    headers = [
        "Projeto", "Hash do Teste", 
        "Média de Cobertura - PASSED (%)", "Desvio Padrão - PASSED", 
        "Média de Cobertura - FAILED (%)", "Desvio Padrão - FAILED"
    ]

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    for row_num, dados in enumerate(resultados, start=2):
        linha = [
            dados['Projeto'], dados['Hash'], 
            dados['AVG_P'], dados['SD_P'], 
            dados['AVG_F'], dados['SD_F']
        ]
        
        for col_num, valor in enumerate(linha, 1):
            cell = ws.cell(row=row_num, column=col_num, value=valor)
            cell.border = thin_border
            if col_num in [1, 2]: 
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else: 
                cell.alignment = center_alignment

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 25

    nome_arquivo = 'Estatisticas_Cobertura_Final.xlsx'
    wb.save(nome_arquivo)
    print(f"\nSucesso! Planilha gerada: {nome_arquivo} com desvio padrão rigorosamente validado.")

if __name__ == "__main__":
    processar_dados()
