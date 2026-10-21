import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
import openpyxl

# ===================== CONFIGURAÇÃO (edite aqui se precisar) =====================
BASE_DIR = Path(__file__).resolve().parent
XLSX_PATH = BASE_DIR / "Profiling_preenchido.xlsx"  # planilha de entrada
OUT_DIR = BASE_DIR                                  # pasta onde os PNGs serão salvos
DPI = 150                                           # resolução das imagens
# ==================================================================================

# Ordem de plotagem: (aba, rótulo no eixo x, grupo)
ORDER = [
    ("ncalls_full", "Ncalls", "count"),
    ("ncalls_top20", "Ncalls_top20", "count"),
    ("tottime_full", "Tottime", "time"),
    ("tottime_top20", "Tottime_top20", "time"),
    ("percall_full", "Percall", "percall"),
    ("percall_top20", "Percall_top20", "percall"),
    ("cumtime_full", "Cumtime", "time"),
    ("cumtime_top20", "Cumtime_top20", "time"),
    ("percall_cum_full", "Percall_cum", "percall"),
    ("percall_cum_top20", "Percall_cum_top20", "percall"),
]

COLORS = {"count": "#99CBCB", "time": "#F4AFAB", "percall": "#F7D68F"}

# Colunas (0-indexadas) nas abas: Result, CV passed, CV failed
COL_RESULT, COL_CV_PASSED, COL_CV_FAILED = 5, 8, 11
FIRST_DATA_ROW = 4  # linhas 1-3 são cabeçalho

TEXTS = {
    "en": {
        "ylabel": "Coefficient of Variation, CV = SD/AVG × 100 (%) — log$_{10}$ scale",
        "title": None,
        "legend": {
            "count": "Call-count metric (ncalls)",
            "time": "Time-based metrics (tottime, cumtime)",
            "percall": "Per-call derived metrics (percall, percall_cum)",
            "line": "CV = 1%",
        },
        "figsize": (16.9, 8.3),
        "suffix": "",
    },
    "pt": {
        "ylabel": "Coeficiente de Variação, CV = SD/AVG × 100 (%) — escala log$_{10}$",
        "title": "Distribuição do Coeficiente de Variação entre os Atributos de Profiling\n"
                 "({n} Testes Flaky Observados; Execuções com Sucesso e Falha Agrupadas)",
        "legend": {
            "count": "Métrica de contagem de chamadas (ncalls)",
            "time": "Métricas baseadas em tempo (tottime, cumtime)",
            "percall": "Métricas derivadas por chamada (percall, percall_cum)",
            "line": "CV = 1%",
        },
        "figsize": (20.0, 10.7),
        "suffix": "_pt",
    },
}


def read_log_cv(ws):
    """Retorna (lista de log10(CV), nº de testes flaky observados) de uma aba."""
    values, n_tests = [], 0
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        result = row[COL_RESULT] if len(row) > COL_RESULT else None
        if not result or str(result).startswith("Non"):
            continue  # só "Observed Flaky Test"
        n_tests += 1
        for col in (COL_CV_PASSED, COL_CV_FAILED):
            v = row[col] if len(row) > col else None
            if isinstance(v, (int, float)) and v > 0:  # 'N/A' e zeros não entram no log
                values.append(math.log10(v))
    return values, n_tests


def load_data(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    data, n_tests = [], 0
    for sheet, label, group in ORDER:
        if sheet not in wb.sheetnames:
            raise SystemExit(f"Aba '{sheet}' não encontrada em {xlsx_path}")
        vals, n = read_log_cv(wb[sheet])
        data.append((label, group, vals))
        n_tests = max(n_tests, n)
    return data, n_tests


def plot(data, n_tests, lang, out_path, dpi):
    t = TEXTS[lang]
    fig, ax = plt.subplots(figsize=t["figsize"])

    positions = list(range(len(data)))
    parts = ax.violinplot(
        [d[2] for d in data], positions=positions, widths=0.8,
        showmedians=True, showextrema=True,
    )
    for body, (_, group, _) in zip(parts["bodies"], data):
        body.set_facecolor(COLORS[group])
        body.set_edgecolor("#555555")
        body.set_linewidth(1.0)
        body.set_alpha(1.0)
    for key in ("cmedians", "cmins", "cmaxes", "cbars"):
        parts[key].set_color("#222222")
        parts[key].set_linewidth(1.6 if key == "cmedians" else 1.3)

    ax.axhline(0, color="gray", linestyle="--", linewidth=1.0)  # CV = 1% -> log10 = 0
    ax.set_xticks(positions)
    ax.set_xticklabels([d[0] for d in data], rotation=15, ha="right", fontsize=11)
    ax.set_ylabel(t["ylabel"], fontsize=12)
    ax.set_ylim(-6.75, 2.75)
    ax.yaxis.grid(True, color="#e8e8e8", linewidth=0.8)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    if t["title"]:
        ax.set_title(t["title"].format(n=n_tests), fontsize=13)

    lg = t["legend"]
    handles = [
        Patch(facecolor=COLORS[g], edgecolor="#555555", label=lg[g])
        for g in ("count", "time", "percall")
    ] + [Line2D([0], [0], color="gray", linestyle="--", linewidth=1, label=lg["line"])]
    ax.legend(handles=handles, loc="lower right", fontsize=10, frameon=True)

    fig.tight_layout()
    fig.savefig(out_path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)
    print(f"Salvo: {out_path}")


def main():
    if not XLSX_PATH.exists():
        raise SystemExit(f"Planilha não encontrada: {XLSX_PATH}\n"
                         "Coloque o script na mesma pasta do Profiling_preenchido.xlsx "
                         "ou ajuste XLSX_PATH.")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data, n_tests = load_data(XLSX_PATH)
    print(f"{n_tests} testes flaky observados")
    for lang, t in TEXTS.items():
        plot(data, n_tests, lang, OUT_DIR / f"cv_violin_reordered_v2{t['suffix']}.png", DPI)


if __name__ == "__main__":
    main()
